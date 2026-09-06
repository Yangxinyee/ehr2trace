"""Draw the end-to-end framework figure that places ehr2cdm among the components a
patient world model or clinical agent system needs.

Writes docs/figures/system_framework.pptx. Every mark is a native PowerPoint shape:
cards and rules are autoshapes, icons are vector custom geometry, and all text is
live text, so the deck can be edited without returning to this script.

Stage 2 is the system this repository implements. Stages 3 to 5 are drawn dashed
because they are requirements the paper identifies, not capabilities it delivers.
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs' / 'figures' / 'system_framework.pptx'

# Palette shared with the manuscript figures (paper/figures/figstyle.tex).
TEAL = RGBColor(0x0D, 0x6E, 0x66)
BLUE = RGBColor(0x1F, 0x56, 0x8C)
ORANGE = RGBColor(0xBF, 0x57, 0x20)
SLATE = RGBColor(0x46, 0x4E, 0x56)
GREY = RGBColor(0x8A, 0x92, 0x99)
RULE = RGBColor(0xD8, 0xDC, 0xE0)
INK = RGBColor(0x1E, 0x24, 0x2A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = 'Arial'

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.44)
COLS, GAP = 5, Inches(0.26)
CARD_W = int((SLIDE_W - 2 * MARGIN - (COLS - 1) * GAP) / COLS)
CARD_T, CARD_H = Inches(1.30), Inches(4.06)


def col_x(i):
    return int(MARGIN + i * (CARD_W + GAP))


def flat(shape):
    """Drop the theme style reference and any inherited shadow.

    A new autoshape carries a <p:style> pointing at the theme's effect list, which
    is where the default drop shadow comes from. Explicit fill and line are set on
    every shape here, so the reference has nothing left to contribute.
    """
    shape.shadow.inherit = False
    style = shape._element.find(qn('p:style'))
    if style is not None:
        shape._element.remove(style)


def box(shapes, kind, x, y, w, h, fill=None, line=None, width=Pt(0.75), dash=None):
    shape = shapes.add_shape(kind, int(x), int(y), int(w), int(h))
    flat(shape)
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = width
        if dash:
            shape.line.dash_style = dash
    shape.text_frame.text = ''
    return shape


def text(shapes, x, y, w, h, runs, size=10, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.0, space_after=2):
    """runs: a string, or a list of (string, {overrides}) for mixed formatting."""
    tb = shapes.add_textbox(int(x), int(y), int(w), int(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = runs if isinstance(runs, list) else [runs]
    for i, line in enumerate(lines):
        body, over = (line, {}) if isinstance(line, str) else line
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = over.get('align', align)
        p.line_spacing = over.get('spacing', spacing)
        p.space_after = Pt(over.get('space_after', space_after))
        emit(p, body, over.get('size', size), over.get('color', color),
             over.get('bold', bold), over.get('italic', False))
    return tb


# --------------------------------------------------------------------- icons
A_NS = 'http://schemas.openxmlformats.org/drawingml/2006/main'


def _sub(parent, tag, **attrs):
    el = etree.SubElement(parent, qn(tag))
    for key, value in attrs.items():
        el.set(key, str(int(value)))
    return el


def _tokens(path_d):
    out, buf = [], ''
    for ch in path_d.replace(',', ' '):
        if ch in 'MLCZ':
            if buf.strip():
                out.extend(buf.split())
            buf = ''
            out.append(ch)
        else:
            buf += ch
    if buf.strip():
        out.extend(buf.split())
    return out


def _custgeom(shape, path_d, span=1000):
    """Replace a shape's preset geometry with an absolute path on a 0..span box."""
    spPr = shape._element.spPr
    preset = spPr.find(qn('a:prstGeom'))
    index = list(spPr).index(preset)
    spPr.remove(preset)
    geom = etree.Element(qn('a:custGeom'))
    for tag in ('a:avLst', 'a:gdLst', 'a:ahLst', 'a:cxnLst'):
        etree.SubElement(geom, qn(tag))
    rect = etree.SubElement(geom, qn('a:rect'))
    for key, value in (('l', 0), ('t', 0), ('r', span), ('b', span)):
        rect.set(key, str(value))
    path = _sub(etree.SubElement(geom, qn('a:pathLst')), 'a:path', w=span, h=span)
    tokens, i = _tokens(path_d), 0
    while i < len(tokens):
        cmd = tokens[i]
        if cmd == 'Z':
            etree.SubElement(path, qn('a:close'))
            i += 1
            continue
        count = {'M': 1, 'L': 1, 'C': 3}[cmd]
        node = etree.SubElement(path, qn({'M': 'a:moveTo', 'L': 'a:lnTo',
                                          'C': 'a:cubicBezTo'}[cmd]))
        for j in range(count):
            _sub(node, 'a:pt', x=float(tokens[i + 1 + 2 * j]),
                 y=float(tokens[i + 2 + 2 * j]))
        i += 1 + 2 * count
    spPr.insert(index, geom)


def icon(shapes, x, y, size, layers, name='icon'):
    """layers: (path, 'stroke'|'fill', colour, stroke width in points)."""
    parts = []
    for path_d, mode, colour, width in layers:
        shape = shapes.add_shape(MSO_SHAPE.RECTANGLE, int(x), int(y), int(size), int(size))
        flat(shape)
        _custgeom(shape, path_d)
        if mode == 'fill':
            shape.fill.solid()
            shape.fill.fore_color.rgb = colour
            shape.line.fill.background()
        else:
            shape.fill.background()
            shape.line.color.rgb = colour
            shape.line.width = Pt(width)
        shape.text_frame.text = ''
        parts.append(shape)
    if len(parts) > 1:
        group = shapes.add_group_shape(parts)
        group.name = name
        return group
    parts[0].name = name
    return parts[0]


def circ(cx, cy, r):
    k = r * 0.5523
    return (f'M {cx - r} {cy} C {cx - r} {cy - k} {cx - k} {cy - r} {cx} {cy - r} '
            f'C {cx + k} {cy - r} {cx + r} {cy - k} {cx + r} {cy} '
            f'C {cx + r} {cy + k} {cx + k} {cy + r} {cx} {cy + r} '
            f'C {cx - k} {cy + r} {cx - r} {cy + k} {cx - r} {cy} Z')


def bar(x, y, w, h):
    return f'M {x} {y} L {x + w} {y} L {x + w} {y + h} L {x} {y + h} Z'


def _doc(dx, dy):
    return (f'M {80 + dx} {140 + dy} L {470 + dx} {140 + dy} L {630 + dx} {300 + dy} '
            f'L {630 + dx} {820 + dy} L {80 + dx} {820 + dy} Z '
            f'M {470 + dx} {140 + dy} L {470 + dx} {300 + dy} L {630 + dx} {300 + dy}')


def icon_sources(colour):
    return [(_doc(280, 0), 'stroke', colour, 1.4),
            (_doc(150, 70), 'stroke', colour, 1.4),
            (_doc(0, 140), 'stroke', colour, 1.8),
            ('M 155 700 L 470 700 M 155 800 L 400 800', 'stroke', colour, 1.4)]


def icon_contract(colour):
    return [('M 500 70 L 890 215 L 890 530 C 890 760 715 905 500 962 '
             'C 285 905 110 760 110 530 L 110 215 Z', 'stroke', colour, 1.9),
            ('M 318 508 L 452 642 L 702 372', 'stroke', colour, 2.4)]


def icon_trajectory(colour):
    """Observed history, the decision time, and the futures conditioned on an action."""
    observed = ' '.join(circ(cx, 620, 58) for cx in (90, 290, 490))
    decision = ' '.join(f'M 640 {y} L 640 {y + 110}' for y in (150, 330, 510, 690, 870))
    futures = circ(880, 370, 78) + ' ' + circ(880, 860, 78)
    return [('M 40 620 L 570 620', 'stroke', colour, 2.0),
            (observed, 'fill', colour, 0),
            (decision, 'stroke', GREY, 1.8),
            ('M 665 620 L 810 435 M 665 620 L 810 800', 'stroke', colour, 1.8),
            (futures, 'stroke', colour, 1.8)]


def icon_learning(colour):
    nodes = [(160, 280), (160, 520), (160, 760),
             (500, 200), (500, 470), (500, 740),
             (850, 330), (850, 640)]
    edges = ['M 160 280 L 500 200', 'M 160 280 L 500 470', 'M 160 520 L 500 200',
             'M 160 520 L 500 740', 'M 160 760 L 500 470', 'M 160 760 L 500 740',
             'M 500 200 L 850 330', 'M 500 470 L 850 330', 'M 500 470 L 850 640',
             'M 500 740 L 850 640']
    return [(' '.join(edges), 'stroke', GREY, 1.2),
            (' '.join(circ(cx, cy, 62) for cx, cy in nodes), 'fill', colour, 0)]


def icon_agent(colour):
    return [('M 500 120 L 500 285', 'stroke', colour, 1.7),
            (circ(500, 96, 52), 'fill', colour, 0),
            ('M 265 285 L 735 285 C 790 285 830 325 830 380 L 830 700 '
             'C 830 755 790 795 735 795 L 265 795 C 210 795 170 755 170 700 '
             'L 170 380 C 170 325 210 285 265 285 Z', 'stroke', colour, 1.8),
            (circ(378, 490, 54) + ' ' + circ(622, 490, 54), 'fill', colour, 0),
            ('M 392 664 L 608 664', 'stroke', colour, 1.7)]


def icon_evaluation(colour):
    bars = ' '.join(bar(x, y, 118, 862 - y) for x, y in ((100, 620), (270, 452), (440, 700)))
    return [('M 70 880 L 640 880', 'stroke', GREY, 1.4),
            (bars, 'fill', colour, 0),
            (circ(690, 372, 196), 'stroke', colour, 2.0),
            ('M 830 512 L 938 620', 'stroke', colour, 2.4)]


# ------------------------------------------------------------------- content
TAG_INPUT, TAG_WORK, TAG_FUTURE = 'input', 'work', 'future'

STAGES = [
    dict(letter='a', title='Heterogeneous EHR sources',
         role='What a hospital actually exports',
         colour=SLATE, tag=TAG_INPUT, tag_text='EXTERNAL INPUT', icon=icon_sources,
         bullets=['CSV, Excel and Parquet extracts',
                  'Orders, administrations, labs and notes',
                  'Registry and imaging cohorts',
                  'Site-local schemas, codes and time zones']),
    dict(letter='b', title='Auditable data infrastructure',
         role='Source-linked conversion under an executable contract',
         colour=TEAL, tag=TAG_WORK, tag_text='ehr2cdm  —  THIS WORK',
         icon=icon_contract,
         bullets=['Canonical events with row-level lineage',
                  'Occurrence time and availability time kept apart',
                  'OMOP CDM 5.4 and MEDS exports',
                  '36 executable checks across stored artifacts',
                  'Quarantine, review records, trace to source row']),
    dict(letter='c', title='Decision-time trajectory layer',
         role='Turning an event stream into learnable transitions',
         colour=ORANGE, ink=SLATE, tag=TAG_FUTURE, tag_text='NOT IMPLEMENTED HERE',
         icon=icon_trajectory,
         bullets=['History H~t~ restricted to what is known by t',
                  'Transitions (o~t~, a~t~, Δt, o~t+1~, r~t~, d~t~, c~t~)',
                  'Episode boundaries and missingness masks',
                  'Action eligibility and reward versioning',
                  'Cohort and task definitions']),
    dict(letter='d', title='World models and agent policies',
         role='What the trajectories are for',
         colour=ORANGE, ink=SLATE, tag=TAG_FUTURE, tag_text='NOT IMPLEMENTED HERE',
         icon=icon_learning, second_icon=icon_agent,
         sections=[('Patient world model',
                    'Action-conditioned prediction p~θ~(o~t+1~, Δt | H~t~, a~t~); '
                    'simulated patient trajectories'),
                   ('Clinical agent',
                    'Offline reinforcement learning; retrieval and tool use; '
                    'conversational patient agents')]),
    dict(letter='e', title='Evaluation and deployment',
         role='Evidence that a model is safe to act on',
         colour=ORANGE, ink=SLATE, tag=TAG_FUTURE, tag_text='NOT IMPLEMENTED HERE',
         icon=icon_evaluation,
         bullets=['Held-out and prospective validation',
                  'Off-policy evaluation under confounding',
                  'Leakage, safety and subgroup audits',
                  'Clinician-in-the-loop review',
                  'Site-to-site portability']),
]

FOUNDATION = ['Time semantics', 'Action semantics', 'Identity resolution',
              'Reviewed terminology']


def emit(paragraph, body, size, colour, bold=False, italic=False):
    """Write body into a paragraph, reading ~x~ as a real subscript run.

    Unicode subscript characters are not in Arial, so a deck using them shows
    boxes on a machine that has the real font rather than a substitute.
    """
    for i, chunk in enumerate(body.split('~')):
        if not chunk:
            continue
        run = paragraph.add_run()
        run.text = chunk
        run.font.name, run.font.size = FONT, Pt(size)
        run.font.bold, run.font.italic = bold, italic
        run.font.color.rgb = colour
        if i % 2:
            run.font._rPr.set('baseline', '-25000')
    return paragraph


def bullets(shapes, x, y, w, items, colour, size=8.5, gap=4.5):
    tb = shapes.add_textbox(int(x), int(y), int(w), Inches(2.4))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.16
        p.space_after = Pt(gap)
        dot = p.add_run()
        dot.text = '•  '
        dot.font.name, dot.font.size, dot.font.bold = FONT, Pt(size), True
        dot.font.color.rgb = colour
        emit(p, item, size, INK)
    return tb


def tag(shapes, x, y, w, h, kind, label, colour):
    if kind == TAG_WORK:
        shape = box(shapes, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=colour)
        ink, bold = WHITE, True
    elif kind == TAG_FUTURE:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        shape = box(shapes, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=WHITE,
                    line=colour, width=Pt(0.9), dash=MSO_LINE_DASH_STYLE.DASH)
        ink, bold = colour, False
    else:
        shape = box(shapes, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=WHITE,
                    line=RULE, width=Pt(0.9))
        ink, bold = GREY, False
    shape.adjustments[0] = 0.5
    text(shapes, x, y + Emu(int(h * 0.20)), w, h, label, size=7,
         color=ink, bold=bold, align=PP_ALIGN.CENTER)


def arrowhead(shape, colour, width=Pt(1.1)):
    """Put a triangular head on a straight connector."""
    shape.line.color.rgb = colour
    shape.line.width = width
    ln = shape.line._get_or_add_ln()
    head = etree.SubElement(ln, qn('a:headEnd'))
    head.set('type', 'triangle')
    head.set('w', 'med')
    head.set('h', 'med')


def feedback(shapes):
    """The loop that closes the figure: what evaluation finds becomes a new check."""
    from pptx.enum.shapes import MSO_CONNECTOR
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    y = CARD_T + CARD_H + Inches(0.42)
    x_from = col_x(4) + CARD_W // 2
    x_to = col_x(1) + CARD_W // 2
    for x in (x_to, x_from):
        drop = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x, CARD_T + CARD_H, x, y)
        drop.line.color.rgb = RULE
        drop.line.width = Pt(1.0)
        drop.name = 'feedback-stem'
    line = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x_to, y, x_from, y)
    arrowhead(line, GREY)
    line.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    line.name = 'feedback'
    text(shapes, x_to, y - Inches(0.26), x_from - x_to, Inches(0.20),
         'Failures found downstream become new contract checks',
         size=8.5, color=SLATE, align=PP_ALIGN.CENTER)


def ribbon(shapes, inner):
    top, height = Inches(6.14), Inches(0.72)
    card = box(shapes, MSO_SHAPE.ROUNDED_RECTANGLE, MARGIN, top, inner, height,
               fill=RGBColor(0xEE, 0xF5, 0xF4))
    card.adjustments[0] = 0.10
    card.name = 'foundation'
    accent = box(shapes, MSO_SHAPE.RECTANGLE, MARGIN, top, Inches(0.045), height,
                 fill=TEAL)
    accent.name = 'foundation-accent'

    left = MARGIN + Inches(0.22)
    text(shapes, left, top + Inches(0.13), Inches(5.4), Inches(0.22),
         'What stage b establishes and every stage to its right depends on',
         size=9, bold=True, color=TEAL)
    text(shapes, left, top + Inches(0.36), Inches(5.4), Inches(0.24),
         'Every downstream claim traces to a source file and row.',
         size=8.5, color=SLATE)

    chip_w, chip_gap = Inches(1.48), Inches(0.12)
    span = 4 * chip_w + 3 * chip_gap
    x = MARGIN + inner - Inches(0.22) - span
    for label in FOUNDATION:
        chip = box(shapes, MSO_SHAPE.ROUNDED_RECTANGLE, x, top + Inches(0.20),
                   chip_w, Inches(0.32), fill=WHITE, line=RGBColor(0xC5, 0xDB, 0xD8),
                   width=Pt(0.8))
        chip.adjustments[0] = 0.35
        chip.name = f'foundation-{label}'
        text(shapes, x, top + Inches(0.28), chip_w, Inches(0.20), label,
             size=8, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
        x += chip_w + chip_gap


def legend(shapes, inner):
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    y, size = Inches(7.02), Inches(0.13)
    items = [(TEAL, None, 'Implemented and evaluated in this paper'),
             (WHITE, ORANGE, 'Required downstream component, identified but not implemented here'),
             (WHITE, RULE, 'External input to the system')]
    x = MARGIN
    for fill, line, label in items:
        key = box(shapes, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, size, size,
                  fill=fill, line=line, width=Pt(1.0),
                  dash=MSO_LINE_DASH_STYLE.DASH if line is ORANGE else None)
        key.adjustments[0] = 0.25
        key.name = 'legend-key'
        text(shapes, x + Inches(0.20), y - Inches(0.01), Inches(4.4), Inches(0.20),
             label, size=8, color=SLATE)
        x += Inches(0.20) + Inches(0.062) * len(label) + Inches(0.34)


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shapes = slide.shapes
    box(shapes, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H, fill=WHITE)

    inner = SLIDE_W - 2 * MARGIN
    text(shapes, MARGIN, Inches(0.30), inner, Inches(0.34),
         'From hospital exports to patient world models and clinical agents',
         size=19, bold=True, color=INK)
    text(shapes, MARGIN, Inches(0.66), inner, Inches(0.26),
         [('ehr2cdm supplies the audited data foundation (b). Stages c to e are '
           'requirements this paper identifies and does not implement.', {})],
         size=10, color=SLATE)

    tag_t, tag_h = Inches(1.00), Inches(0.235)
    from pptx.enum.dml import MSO_LINE_DASH_STYLE

    for i, stage in enumerate(STAGES):
        x, colour = col_x(i), stage['colour']
        ink = stage.get('ink', colour)
        tag(shapes, x + Emu(int(CARD_W * 0.06)), tag_t, int(CARD_W * 0.88), tag_h,
            stage['tag'], stage['tag_text'], colour)

        work = stage['tag'] == TAG_WORK
        card = box(shapes, MSO_SHAPE.ROUNDED_RECTANGLE, x, CARD_T, CARD_W, CARD_H,
                   fill=RGBColor(0xF3, 0xF8, 0xF7) if work else WHITE,
                   line=colour if stage['tag'] != TAG_INPUT else RULE,
                   width=Pt(1.9) if work else Pt(1.0),
                   dash=None if stage['tag'] != TAG_FUTURE else MSO_LINE_DASH_STYLE.DASH)
        card.adjustments[0] = 0.035
        card.name = f"stage-{stage['letter']}"

        pad = Inches(0.17)
        cx, cw = x + pad, CARD_W - 2 * pad
        text(shapes, cx, CARD_T + Inches(0.15), cw, Inches(0.30),
             [(stage['letter'], {'size': 13, 'bold': True, 'color': colour})])
        text(shapes, cx + Inches(0.24), CARD_T + Inches(0.155), cw - Inches(0.24),
             Inches(0.46), stage['title'], size=10.5, bold=True, color=INK, spacing=1.08)
        text(shapes, cx, CARD_T + Inches(0.66), cw, Inches(0.30), stage['role'],
             size=8, color=GREY, spacing=1.12)

        rule = box(shapes, MSO_SHAPE.RECTANGLE, cx, CARD_T + Inches(1.00),
                   cw, Emu(9525), fill=RULE)
        rule.name = 'rule'

        size = Inches(0.62)
        if 'second_icon' in stage:
            icon(shapes, cx + Inches(0.10), CARD_T + Inches(1.16), size,
                 stage['icon'](ink), 'icon-world-model')
            icon(shapes, cx + cw - size - Inches(0.10), CARD_T + Inches(1.16), size,
                 stage['second_icon'](ink), 'icon-agent')
        else:
            icon(shapes, x + (CARD_W - size) / 2, CARD_T + Inches(1.16), size,
                 stage['icon'](ink), f"icon-{stage['letter']}")

        top = CARD_T + Inches(1.92)
        if 'sections' in stage:
            for head, body in stage['sections']:
                text(shapes, cx, top, cw, Inches(0.22), head, size=8.5, bold=True,
                     color=colour)
                text(shapes, cx, top + Inches(0.20), cw, Inches(0.70), body, size=8.5,
                     color=INK, spacing=1.16)
                top += Inches(0.92)
        else:
            bullets(shapes, cx, top, cw, stage['bullets'], colour if colour is not ORANGE
                    else RGBColor(0xC9, 0xA0, 0x8A))

        if i < len(STAGES) - 1:
            arrow = box(shapes, MSO_SHAPE.ISOSCELES_TRIANGLE,
                        x + CARD_W + Inches(0.05), CARD_T + Inches(1.28),
                        Inches(0.16), Inches(0.19), fill=RGBColor(0xAE, 0xB6, 0xBD))
            arrow.rotation = 90
            arrow.name = 'flow'

    feedback(shapes)
    ribbon(shapes, inner)
    legend(shapes, inner)
    return prs, shapes, slide


def export_pdf():
    """Vector PDF alongside the deck, when LibreOffice is available."""
    import shutil
    import subprocess
    soffice = shutil.which('soffice') or shutil.which('libreoffice')
    if not soffice:
        print('No LibreOffice on PATH; skipped the PDF export.')
        return None
    subprocess.run([soffice, '--headless', '--convert-to', 'pdf',
                    '--outdir', str(OUT.parent), str(OUT)],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    return OUT.with_suffix('.pdf')


if __name__ == '__main__':
    prs, _, _ = build()
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f'Wrote {OUT.relative_to(ROOT)}')
    pdf = export_pdf()
    if pdf:
        print(f'Wrote {pdf.relative_to(ROOT)}')
