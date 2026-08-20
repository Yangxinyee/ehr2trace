# -*- coding: utf-8 -*-
"""
verify_doc_baselines.py — re-derive every number in design spec section 2 from the raw
data and compare it, item by item, against what the document asserts.

The design spec claims all input facts were empirically verified. This script makes that
claim checkable: if someone edits a number in the document, or the data itself changes,
one run finds it.

Usage:
    python3 tools/verify_doc_baselines.py            # all checks (about one minute)
    python3 tools/verify_doc_baselines.py --fast     # skip the full-table scans

Exit code 0 means everything agrees. Non-zero means drift, and drift should be reported
and investigated in the data rather than papered over by editing the document.
"""
from __future__ import annotations

import argparse
import json
import os
import subprocess
from pathlib import Path

ROOT = Path(os.environ.get("EHR_DATA_ROOT", ""))

PARTITIONS = {
    "29_has":  "29 - pulmonary embolism/has_embolism",
    "29_no":   "29 - pulmonary embolism/no_embolism",
    "29b_has": "29b - pulmonary embolism/has_pulmonary_embolism",
    "29b_no":  "29b - pulmonary embolism/no_pulmonary_embolism",
}
XLSX = {
    "29_has":  "29_has_pulmonary_embolism.xlsx",
    "29_no":   "29_no_embolism.xlsx",
    "29b_has": "29b_has_pulmonary_embolism.xlsx",
    "29b_no":  "29b_no_pulmonary_embolism.xlsx",
}

# The pseudonym -> real MRN mapping and the patient-level dates live in a local file that
# is never committed. See the data sensitivity section of the README.
LOCAL_BASELINES = Path(__file__).with_name("baselines.local.json")


def load_local() -> dict:
    if not LOCAL_BASELINES.exists():
        print(f"Missing local baseline file: {LOCAL_BASELINES}\n"
              f"It holds real MRNs and patient-level dates and is git-ignored by design.\n"
              f"Generate it with:\n"
              f"    python3 tools/make_local_baselines.py PT-A=<MRN> PT-B=<MRN> PT-C=<MRN>\n")
        raise SystemExit(2)
    return json.loads(LOCAL_BASELINES.read_text(encoding="utf-8"))


# ---- Expected values asserted by design spec section 2 ---------------------
EXPECT = {
    "n_physical_files": 25,
    "n_txt_sources": 21,
    "n_xlsx_sheets": 19,
    "mrn_per_partition": {"29_has": 5633, "29_no": 8623, "29b_has": 8247, "29b_no": 8468},
    "mrn_union": 22982,
    "pairwise": {("29_has", "29_no"): 723, ("29_has", "29b_has"): 5633,
                 ("29_has", "29b_no"): 392, ("29_no", "29b_has"): 923,
                 ("29_no", "29b_no"): 747, ("29b_has", "29b_no"): 927},
    "has_no_overlap": 1609,
    "29has_subset_of_29bhas": True,
    "sheet_names": {
        "29_has":  ["Demographics", "Medication Administration", "PFT Narrative", "PFT Values", "Outcome"],
        "29_no":   ["Demographics", "Medication Administration", "PFT", "PFT Value", "Outcome"],
        "29b_has": ["Demographics", "PFT Narrative", "PFT Values", "Outcome"],
        "29b_no":  ["Demographics", "Medication Administration", "PFT Narrative", "PFT Values", "Outcome"],
    },
    # Per-partition row counts for the three patients (design section 2.5)
    "patient_rows": {
        "29_has":  {"PT-A": (5727, 379), "PT-B": (4243, 206), "PT-C": (377, 22)},
        "29b_has": {"PT-A": (6095, 11),  "PT-B": (6162, 1),   "PT-C": (398, 1)},
        "29b_no":  {"PT-B": (2529, 206)},
    },
    "patient_partitions": {
        "PT-A": {"29_has", "29b_has"},
        "PT-B": {"29_has", "29b_has", "29b_no"},
        "PT-C": {"29_has", "29b_has"},
    },
    # PT-B's anchor dates are patient-level service dates, so the expected values live in
    # baselines.local.json rather than in this repository.
    "ptb_n_anchors": 4,
    "ptb_anchors_per_partition": {"29_has": 2, "29b_has": 3, "29b_no": 1},
    "ptb_echo_rows_per_anchor": 547,
    "echo_description_variants": 14,
    "bom_present": 19,
    "bom_absent": ["29_has_embolism_problem_list.txt", "29_no_embolism_problem_list.txt"],
    "type_drift": {  # (partition, sheet, column) -> expected Python type name
        ("29_has", "Demographics", "Death_Date"): "datetime",
        ("29b_has", "Demographics", "Death_Date"): "str",
        ("29_no", "Outcome", "Length_of_stay_days"): "str",
        ("29_has", "Outcome", "Length_of_stay_days"): "int",
    },
}

results: list[tuple[bool, str, str]] = []
LOCAL: dict = {}
ALIAS: dict[str, str] = {}


def check(name: str, got, want) -> None:
    ok = got == want
    results.append((ok, name, "" if ok else f"expected {want!r}, measured {got!r}"))
    print(f"  {'OK  ' if ok else 'FAIL'}  {name}"
          + ("" if ok else f"   <- expected {want!r}, measured {got!r}"))


# ---------------------------------------------------------------- checks
def check_files() -> None:
    print("\n[1] File and logical source inventory")
    txts, xlsxs = [], []
    for pid, rel in PARTITIONS.items():
        d = ROOT / rel
        txts += sorted(d.glob("*.txt"))
        xlsxs += sorted(d.glob("*.xlsx"))
    check("physical file count", len(txts) + len(xlsxs), EXPECT["n_physical_files"])
    check("txt logical source count", len(txts), EXPECT["n_txt_sources"])

    import openpyxl
    total_sheets = 0
    for pid, rel in PARTITIONS.items():
        wb = openpyxl.load_workbook(ROOT / rel / XLSX[pid], read_only=True)
        check(f"{pid} sheet names", wb.sheetnames, EXPECT["sheet_names"][pid])
        total_sheets += len(wb.sheetnames)
        wb.close()
    check("total xlsx sheets", total_sheets, EXPECT["n_xlsx_sheets"])


def check_bom() -> None:
    """Design section 2.1: BOM presence is inconsistent -- 19 txt files have one, 2 do not."""
    print("\n[2] UTF-8 BOM distribution")
    with_bom, without = [], []
    for pid, rel in PARTITIONS.items():
        for f in sorted((ROOT / rel).glob("*.txt")):
            (with_bom if f.open("rb").read(3) == b"\xef\xbb\xbf" else without).append(f.name)
    check("txt files with a BOM", len(with_bom), EXPECT["bom_present"])
    check("txt files without a BOM", sorted(without), EXPECT["bom_absent"])


def load_mrns() -> dict[str, set[str]]:
    import openpyxl
    out = {}
    for pid, rel in PARTITIONS.items():
        wb = openpyxl.load_workbook(ROOT / rel / XLSX[pid], read_only=True)
        rows = wb["Demographics"].iter_rows(values_only=True)
        next(rows)
        out[pid] = {str(r[0]).strip() for r in rows if r and r[0]}
        wb.close()
    return out


def check_cohort(mrns: dict[str, set[str]]) -> None:
    print("\n[3] Partition and patient overlap baselines")
    for pid, n in EXPECT["mrn_per_partition"].items():
        check(f"{pid} unique MRNs", len(mrns[pid]), n)
    check("union of all four partitions", len(set().union(*mrns.values())), EXPECT["mrn_union"])
    for (a, b), n in EXPECT["pairwise"].items():
        check(f"{a} intersect {b}", len(mrns[a] & mrns[b]), n)
    has = mrns["29_has"] | mrns["29b_has"]
    no = mrns["29_no"] | mrns["29b_no"]
    check("has intersect no", len(has & no), EXPECT["has_no_overlap"])
    check("29_has is a subset of 29b_has",
          mrns["29_has"] <= mrns["29b_has"], EXPECT["29has_subset_of_29bhas"])


def check_patients(mrns: dict[str, set[str]]) -> None:
    print("\n[4] Which partitions each patient appears in")
    for m, want in EXPECT["patient_partitions"].items():
        got = {pid for pid in PARTITIONS if ALIAS[m] in mrns[pid]}
        check(f"{m} appears in", got, want)


def grep_count(path: Path, mrn: str) -> int:
    r = subprocess.run(["grep", "-c", f"^{mrn}\t", str(path)],
                       capture_output=True, text=True)
    return int(r.stdout.strip() or 0)


def check_patient_rows() -> None:
    print("\n[5] Per-partition row counts for the three patients (design section 2.5)")
    import openpyxl
    for pid, per_patient in EXPECT["patient_rows"].items():
        d = ROOT / PARTITIONS[pid]
        wb = openpyxl.load_workbook(d / XLSX[pid], read_only=True)
        # Pre-scan the workbook once, counting across all sheets.
        xcount = {ALIAS[m]: 0 for m in per_patient}
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            next(rows)
            for r in rows:
                if not r:
                    continue
                # 29b_no's PFT Narrative has an empty leading column, so the MRN may be
                # in column 0 or column 1.
                for v in r[:2]:
                    s = str(v).strip() if v is not None else ""
                    if s in xcount:
                        xcount[s] += 1
                        break
        wb.close()
        for m, (want_txt, want_xlsx) in per_patient.items():
            got_txt = sum(grep_count(f, ALIAS[m]) for f in sorted(d.glob("*.txt")))
            check(f"{pid} / {m} txt rows", got_txt, want_txt)
            check(f"{pid} / {m} xlsx rows", xcount[ALIAS[m]], want_xlsx)


def check_anchors() -> None:
    print("\n[6] PT-B's CT anchors (design sections 2.3 and 2.5)")
    anchors, per_anchor = set(), {}
    for pid in ("29_has", "29b_has", "29b_no"):
        f = next((ROOT / PARTITIONS[pid]).glob("*echo.txt"))
        r = subprocess.run(["grep", f"^{ALIAS['PT-B']}\t", str(f)],
                           capture_output=True, text=True)
        counts: dict[str, int] = {}
        for line in r.stdout.splitlines():
            # 29 carries a full timestamp and 29b a bare date; take the first 10
            # characters so both compare at date granularity.
            dos = line.split("\t")[3][:10]
            counts[dos] = counts.get(dos, 0) + 1
        anchors |= set(counts)
        per_anchor[pid] = counts
    check("total anchors", len(anchors), EXPECT["ptb_n_anchors"])
    check("anchor set matches the local baseline", anchors, set(LOCAL["ptb_anchors"]))
    for pid, n in EXPECT["ptb_anchors_per_partition"].items():
        check(f"{pid} anchor count", len(per_anchor[pid]), n)
    every = {n for c in per_anchor.values() for n in c.values()}
    check("echo rows per anchor are constant", every, {EXPECT["ptb_echo_rows_per_anchor"]})


def check_dos_format() -> None:
    print("\n[7] dos format differs between batches (design section 2.3 #2)")
    for pid, want_len in (("29_has", 27), ("29b_has", 10)):
        f = next((ROOT / PARTITIONS[pid]).glob("*labs.txt"))
        with f.open(encoding="utf-8", errors="replace") as fh:
            fh.readline()
            dos = fh.readline().split("\t")[1]
        shape = "full timestamp" if want_len > 10 else "date only"
        check(f"{pid} dos length ({shape})", len(dos), want_len)


def check_types() -> None:
    print("\n[8] Cross-workbook cell type drift (design section 2.3 #15)")
    import openpyxl
    for (pid, sheet, col), want in EXPECT["type_drift"].items():
        wb = openpyxl.load_workbook(ROOT / PARTITIONS[pid] / XLSX[pid], read_only=True)
        ws = wb[sheet]
        rows = ws.iter_rows(max_row=2, values_only=True)
        hdr = [str(c) for c in next(rows)]
        row = next(rows)
        got = type(row[hdr.index(col)]).__name__
        check(f"{pid}.{sheet}.{col} type", got, want)
        wb.close()


def check_empty_first_column() -> None:
    print("\n[9] 29b_no's empty leading column in PFT Narrative (design section 2.2)")
    import openpyxl
    wb = openpyxl.load_workbook(ROOT / PARTITIONS["29b_no"] / XLSX["29b_no"], read_only=True)
    hdr = next(wb["PFT Narrative"].iter_rows(max_row=1, values_only=True))
    check("first column header is empty", hdr[0], None)
    check("second column is mrn", hdr[1], "mrn")
    wb.close()


def check_echo_descriptions() -> None:
    print("\n[10] echo DESCRIPTION values (design section 2.3 #20: no CT/CTPA present)")
    vals: set[str] = set()
    for pid, rel in PARTITIONS.items():
        f = next((ROOT / rel).glob("*echo.txt"))
        r = subprocess.run(f"cut -f3 '{f}' | sort -u", shell=True,
                           capture_output=True, text=True)
        vals |= {v for v in r.stdout.splitlines() if v and v != "DESCRIPTION"}
    check("distinct values", len(vals), EXPECT["echo_description_variants"])
    ct = [v for v in vals if "CT" in v.upper().split() or "CTA" in v.upper()]
    check("CT/CTPA entries among them", ct, [])


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--fast", action="store_true",
                    help="skip the checks that require full-table scans")
    a = ap.parse_args()

    global LOCAL, ALIAS
    LOCAL = load_local()
    ALIAS = LOCAL["alias_to_mrn"]

    print(f"Data root: {ROOT}")
    if not ROOT.exists():
        print("Data root not found. Set EHR_DATA_ROOT and retry.")
        return 2

    check_files()
    check_bom()
    check_empty_first_column()
    check_types()
    check_dos_format()
    mrns = load_mrns()
    check_cohort(mrns)
    check_patients(mrns)
    check_anchors()
    if not a.fast:
        check_patient_rows()
        check_echo_descriptions()

    n_fail = sum(1 for ok, _, _ in results if not ok)
    print(f"\n{'=' * 60}\n{len(results)} checks, {n_fail} failed")
    if n_fail:
        print("\nDrift detected. Report the drift and investigate the data;"
              " do not edit the expected values in the document.")
        for ok, name, msg in results:
            if not ok:
                print(f"  - {name}: {msg}")
    else:
        print("Every assertion in design spec section 2 agrees with the raw data.")
    return 1 if n_fail else 0


if __name__ == "__main__":
    raise SystemExit(main())
