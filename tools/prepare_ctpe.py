"""Project the two `All_kinds/` tables the CTPE conversion reads into its partition layout.

The CTPE export was delivered twice: the four partition directories under
`EHR_DATA_ROOT`, and a second tree, `All_kinds/`, holding tables for several cohorts
side by side. Two of those tables belong to this cohort and are in scope (D-R12,
2026-09-13): the most recent follow-up contact per patient (`followup.xlsx`, one sheet
per partition) and the ADT department stays (`ICU/ICU/*.xlsx`, one workbook per
partition). Everything else under `All_kinds/` is left unread and named in the manifest
with its reason, so the coverage check can account for every delivered file.

Same contract as tools/prepare_cu.py: projection only, no imputation, a sha256 manifest
across the boundary. The output keeps the delivery's own column names and cell text, so
the dataset YAML's declarations refer to what was delivered, plus one constant column on
the follow-up table (`observation_code`) that names the fact the row records, the way
prepare_cu.py projects `visit_type` for an ICU stay.

Which delivery file holds which partition is a lookup verified by MRN in the audit
(P-J4): the file names do not spell it out (`29_pulmonary_embolism` is the *has* group
and `29b_pulmonary_embolism` the *no* group). The output goes under the partition
directories `datasets/ctpe.yaml` declares, beneath a root of its own, so the YAML can
declare both sources with `root_env: CTPE_PREPARED_ROOT`.

Before anything is written the four groups are fingerprinted and compared (remediation
plan T2.J7): column names and order, the type of every cell, the digit-masked shape of
values, null rates, and how many of the partition's own patients the table covers. A
difference that could let a model read the has/no label off the *format* of a row --
a column typed differently in one group, a null pattern only one group has -- stops
the script; the comparison is written to the manifest either way. Only counts, shares
and masked shapes are recorded: no MRN and no cell value leaves this script.

--sample N keeps every Nth patient of each group, by sorted MRN, so a whole-pipeline
run over the prepared tables takes minutes.
"""
from __future__ import annotations

import argparse
import collections
import datetime as dt
import hashlib
import json
import os
import re
import sys
from pathlib import Path
from typing import Any, Iterable, Sequence

import pyarrow as pa
import pyarrow.parquet as pq
import yaml

FOLLOWUP_FILE = "Most Recent Follow-Up/Most Recent Follow-Up/followup.xlsx"
FOLLOWUP_COLUMNS = ["PAT_MRN_ID", "EFFECTIVE_DATE_DTTM", "rn"]
#: the one column this script adds: what a follow-up row is, as a source code
FOLLOWUP_CODE_COLUMN = "observation_code"
FOLLOWUP_CODE = "MOST_RECENT_FOLLOWUP"
ICU_COLUMNS = [
    "PAT_MRN_ID", "ADT_DEPARTMENT_ID", "ADT_DEPARTMENT_NAME", "ADT_LOC_NAME", "IN_DTTM",
    "PAT_OUT_DTTM", "NAME", "ICU_DEPT_YN", "Time_Difference", "PAT_ENC_CSN_ID",
]
MRN_COLUMN = "PAT_MRN_ID"
DATE_COLUMNS = {"followup": ["EFFECTIVE_DATE_DTTM"], "icu_transfers": ["IN_DTTM", "PAT_OUT_DTTM"]}

#: partition id -> where the delivery keeps that partition's rows. Verified by MRN
#: overlap in the audit of 2026-09-13 (P-J4); the names alone would mislead.
GROUPS: dict[str, dict[str, str]] = {
    "29_has": {"followup_sheet": "29_pulmonary_embolism", "icu_file": "ICU/ICU/29_pulmonary_embolism.xlsx"},
    "29_no": {"followup_sheet": "29_no_pulmonary_embolism", "icu_file": "ICU/ICU/29_no_pulmonary_embolism.xlsx"},
    "29b_has": {"followup_sheet": "29b_has_pulmonary_embolism", "icu_file": "ICU/ICU/29b_has_pulmonary_embolism.xlsx"},
    "29b_no": {"followup_sheet": "29b_pulmonary_embolism", "icu_file": "ICU/ICU/29b_pulmonary_embolism.xlsx"},
}

#: the partition's own patient list, read from the raw export for the coverage figure
PARTITION_WORKBOOK_GLOB = "*.xlsx"
PARTITION_MRN_SHEET = "Demographics"
PARTITION_MRN_COLUMN = "MRN"

#: Why every other file under All_kinds/ stays unread (D-R12, P-J12). Matched in order
#: against the path relative to the All_kinds root; the first hit wins.
UNREAD_REASONS: list[tuple[str, str]] = [
    (r"(^|/)__MACOSX/", "AppleDouble resource-fork metadata written by macOS when the archive was made; not data"),
    (r"(^|/)(28|28a|30b|33b|35)_", "another cohort's table (coronary CT, cardiac MRI, cholecystitis ultrasound, ED head CT); not the CTPE cohort"),
    (r"^Surgical Pathology Reports/", "surgical pathology report text: out of scope (D-R12)"),
    (r"^Cardiac Cath/", "cardiac catheterisation reports, covering about 11% of patients: out of scope (D-R12)"),
    (r"^Labs - Lipid Panel/", "lipid panel laboratory values, one workbook for every cohort: out of scope (D-R12)"),
    (r"^ICU/ICU/.*_notes?\.txt$|^ICU/ICU/29b_pulmonary_embolism\.txt$",
     "ICU progress-note text with no header row (note text from byte 0), also holding ED, admission and pre-procedure notes: out of scope (D-R12); never opened"),
    (r"^Surgical Cases/Surgical Cases/29_pulmonary_embolism\.xlsx$",
     "surgical case list: out of scope (D-R12). Mislabelled: same row count as 29_no and its MRNs all belong to 29_no, so it is a copy and the 29_has list is missing from the delivery (P-J12, raised with the data owner)"),
    (r"^Surgical Cases/", "surgical case list, 30-77 rows per group: out of scope (D-R12)"),
    (r"^Ultrasound/", "another cohort's table (cholecystitis ultrasound); not the CTPE cohort"),
]

#: tolerances of the group comparison, as shares of rows
TYPE_OR_SHAPE_PRESENT = 0.01     # a type or shape carried by at least this share of one group ...
NULL_RATE_GAP = 0.05             # ... and absent from another is a hard difference; so is this null-rate gap
COVERAGE_GAP = 0.05
SOFT_GAP = 0.01


# --------------------------------------------------------------------------------
# reading
# --------------------------------------------------------------------------------


def sha256_file(path: Path, chunk: int = 1 << 20) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        while block := fh.read(chunk):
            h.update(block)
    return h.hexdigest()


def read_sheet(path: Path, sheet: str | None = None) -> tuple[list[str], list[tuple]]:
    """Header and data rows of one worksheet, cells as openpyxl typed them.

    Fully empty rows are skipped, as the converter's own workbook adapter skips them.
    """
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
        it = ws.iter_rows(values_only=True)
        header = [str(c).strip() if c is not None else "" for c in next(it)]
        rows = [tuple(r) for r in it if r is not None and not all(v is None for v in r)]
    finally:
        wb.close()
    return header, rows


def sheet_names(path: Path) -> list[str]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def render_cell(value: object) -> str | None:
    """A cell as text, the way the converter's canonical form spells the same cell.

    Text is kept as written -- the delivery's own `NULL` literal included, which the
    dataset YAML declares as a null. Typed cells get the one spelling the converter
    would give them, so a timestamp reads the same whether it came through this script
    or straight from a workbook.
    """
    if value is None:
        return None
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, str):
        return value
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value != value:  # NaN
            return None
        return str(int(value)) if value.is_integer() else repr(value)
    if isinstance(value, dt.datetime):
        return value.replace(microsecond=0, tzinfo=None).isoformat(timespec="seconds")
    if isinstance(value, dt.date):
        return dt.datetime.combine(value, dt.time()).isoformat(timespec="seconds")
    raise TypeError(f"no text form for cell type {type(value)!r}")


def is_null(value: object) -> bool:
    return value is None or (isinstance(value, str) and value.strip() in ("", "NULL"))


def project(header: Sequence[str], rows: Iterable[tuple], columns: Sequence[str]) -> list[list[object]]:
    """Only the named columns, in the named order; a missing column is an error."""
    missing = [c for c in columns if c not in header]
    if missing:
        raise SystemExit(f"columns {missing} not in header {list(header)}")
    idx = [header.index(c) for c in columns]
    return [[row[i] if i < len(row) else None for i in idx] for row in rows]


# --------------------------------------------------------------------------------
# the group fingerprint
# --------------------------------------------------------------------------------


def mask_shape(value: object, width: int = 40) -> str:
    """The shape of a value with every digit and letter replaced; never the value."""
    if value is None:
        return "<none>"
    if isinstance(value, dt.datetime):
        text = value.isoformat(sep=" ", timespec="seconds")
    elif isinstance(value, dt.date):
        text = value.isoformat()
    else:
        text = str(value)
    text = re.sub(r"[0-9]", "9", text)
    text = re.sub(r"[A-Za-z]", "a", text)
    return text[:width]


def fingerprint(
    columns: Sequence[str],
    rows: Sequence[Sequence[object]],
    *,
    partition_mrns: set[str] | None,
    mrn_column: str,
    date_columns: Sequence[str] = (),
    top_shapes: int = 8,
) -> dict[str, Any]:
    """Format statistics of one group's table. Counts, shares and shapes only."""
    n = len(rows)
    out: dict[str, Any] = {"rows": n, "columns": list(columns), "per_column": {}}
    for j, col in enumerate(columns):
        types: collections.Counter = collections.Counter()
        shapes: collections.Counter = collections.Counter()
        years: collections.Counter = collections.Counter()
        distinct: set = set()
        nulls = 0
        for row in rows:
            v = row[j] if j < len(row) else None
            types[type(v).__name__] += 1
            if is_null(v):
                nulls += 1
            shapes[mask_shape(v)] += 1
            distinct.add(render_cell(v) if not isinstance(v, str) else v)
            if col in date_columns and isinstance(v, dt.date):
                years[v.year] += 1
        entry: dict[str, Any] = {
            "types": {k: round(c / n, 6) for k, c in sorted(types.items())} if n else {},
            "null_share": round(nulls / n, 6) if n else 0.0,
            "distinct": len(distinct),
            "shapes": {s: round(c / n, 6) for s, c in shapes.most_common(top_shapes)} if n else {},
        }
        if col in date_columns:
            entry["years"] = {str(y): c for y, c in sorted(years.items())}
        out["per_column"][col] = entry
    if mrn_column in columns:
        k = list(columns).index(mrn_column)
        table_mrns = {str(r[k]).strip() for r in rows if not is_null(r[k])}
        out["distinct_mrns"] = len(table_mrns)
        if partition_mrns is not None:
            out["partition_mrns"] = len(partition_mrns)
            out["coverage_of_partition"] = (
                round(len(table_mrns & partition_mrns) / len(partition_mrns), 6) if partition_mrns else None
            )
            out["share_of_table_mrns_outside_partition"] = (
                round(len(table_mrns - partition_mrns) / len(table_mrns), 6) if table_mrns else None
            )
    return out


def compare(groups: dict[str, dict[str, Any]]) -> dict[str, Any]:
    """Pairwise comparison of the group fingerprints; hard differences stop the write."""
    hard: list[str] = []
    soft: list[str] = []
    names = sorted(groups)
    for i, a in enumerate(names):
        for b in names[i + 1:]:
            fa, fb = groups[a], groups[b]
            if fa["columns"] != fb["columns"]:
                hard.append(f"{a} vs {b}: columns differ ({fa['columns']} vs {fb['columns']})")
                continue
            for col in fa["columns"]:
                ca, cb = fa["per_column"][col], fb["per_column"][col]
                for kind in ("types", "shapes"):
                    keys = set(ca[kind]) | set(cb[kind])
                    for key in sorted(keys):
                        sa, sb = ca[kind].get(key, 0.0), cb[kind].get(key, 0.0)
                        gap = abs(sa - sb)
                        if (sa == 0 or sb == 0) and max(sa, sb) >= TYPE_OR_SHAPE_PRESENT:
                            hard.append(f"{a} vs {b}: {col} {kind[:-1]} {key!r} share {sa} vs {sb}")
                        elif gap >= NULL_RATE_GAP:
                            hard.append(f"{a} vs {b}: {col} {kind[:-1]} {key!r} share {sa} vs {sb}")
                        elif gap >= SOFT_GAP:
                            soft.append(f"{a} vs {b}: {col} {kind[:-1]} {key!r} share {sa} vs {sb}")
                gap = abs(ca["null_share"] - cb["null_share"])
                if gap >= NULL_RATE_GAP:
                    hard.append(f"{a} vs {b}: {col} null share {ca['null_share']} vs {cb['null_share']}")
                elif gap >= SOFT_GAP:
                    soft.append(f"{a} vs {b}: {col} null share {ca['null_share']} vs {cb['null_share']}")
                ya, yb = ca.get("years"), cb.get("years")
                if ya and yb and (min(ya) != min(yb) or max(ya) != max(yb)):
                    soft.append(f"{a} vs {b}: {col} year range {min(ya)}-{max(ya)} vs {min(yb)}-{max(yb)}")
            cov_a, cov_b = fa.get("coverage_of_partition"), fb.get("coverage_of_partition")
            if cov_a is not None and cov_b is not None:
                gap = abs(cov_a - cov_b)
                if gap >= COVERAGE_GAP:
                    hard.append(f"{a} vs {b}: coverage of partition {cov_a} vs {cov_b}")
                elif gap >= SOFT_GAP:
                    soft.append(f"{a} vs {b}: coverage of partition {cov_a} vs {cov_b}")
    return {"verdict": "differs" if hard else "consistent", "hard": hard, "soft": soft}


# --------------------------------------------------------------------------------
# writing
# --------------------------------------------------------------------------------


def write_parquet(columns: Sequence[str], rows: Sequence[Sequence[object]], dest: Path) -> int:
    """All-text parquet, written whole or not at all."""
    dest.parent.mkdir(parents=True, exist_ok=True)
    arrays = {c: pa.array([render_cell(r[j]) for r in rows], type=pa.string()) for j, c in enumerate(columns)}
    table = pa.table(arrays)
    tmp = dest.with_suffix(dest.suffix + ".partial")
    pq.write_table(table, tmp, compression="zstd")
    os.replace(tmp, dest)
    return table.num_rows


def partition_dirs(dataset_yaml: Path) -> dict[str, str]:
    """partition id -> directory, from the dataset YAML, so the layout is declared once."""
    raw = yaml.safe_load(dataset_yaml.read_text(encoding="utf-8"))
    return {p["id"]: p["dir"] for p in raw["partitions"]}


def partition_mrns(raw_root: Path, partition_dir: str) -> tuple[set[str], Path]:
    """The partition's own patients, from its workbook's demographics sheet."""
    pdir = raw_root / partition_dir
    books = sorted(p for p in pdir.glob(PARTITION_WORKBOOK_GLOB) if not p.name.startswith("~$"))
    if len(books) != 1:
        raise SystemExit(f"expected one workbook under {pdir}, found {len(books)}")
    header, rows = read_sheet(books[0], PARTITION_MRN_SHEET)
    k = header.index(PARTITION_MRN_COLUMN)
    return {str(r[k]).strip() for r in rows if not is_null(r[k])}, books[0]


def sample_rows(rows: list[list[object]], mrn_index: int, every: int) -> list[list[object]]:
    """Every Nth patient of the group, by sorted MRN: the same N always keeps the same people."""
    if every <= 1:
        return rows
    mrns = sorted({str(r[mrn_index]).strip() for r in rows if not is_null(r[mrn_index])})
    keep = {m for i, m in enumerate(mrns) if i % every == 0}
    return [r for r in rows if not is_null(r[mrn_index]) and str(r[mrn_index]).strip() in keep]


def unread_inputs(all_kinds: Path, read: set[Path], followup_read_sheets: set[str]) -> list[dict[str, str]]:
    """Every file under the delivery this script did not read, with its reason."""
    out: list[dict[str, str]] = []
    for path in sorted(p for p in all_kinds.rglob("*") if p.is_file()):
        rel = str(path.relative_to(all_kinds))
        if path in read:
            if path.name == Path(FOLLOWUP_FILE).name:
                for sheet in sheet_names(path):
                    if sheet not in followup_read_sheets:
                        out.append({"path": f"{rel}::{sheet}", "reason": "sheet for another cohort; not the CTPE cohort"})
            continue
        reason = next((why for pattern, why in UNREAD_REASONS if re.search(pattern, rel)), None)
        out.append({"path": rel, "reason": reason or "UNCLASSIFIED: a file this script does not know; decide and name it"})
    return out


# --------------------------------------------------------------------------------


def main(argv: Sequence[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--all-kinds-root", required=True, type=Path, help="the All_kinds delivery")
    ap.add_argument("--raw-root", type=Path, default=None,
                    help="the raw export (partition workbooks give the patient lists); default $EHR_DATA_ROOT")
    ap.add_argument("--out", required=True, type=Path, help="the prepared root (CTPE_PREPARED_ROOT)")
    ap.add_argument("--dataset-yaml", type=Path, default=Path(__file__).resolve().parents[1] / "datasets" / "ctpe.yaml")
    ap.add_argument("--sample", type=int, default=0, help="keep 1 patient in N per group (0 = all)")
    ap.add_argument("--write-despite-difference", metavar="REASON", default=None,
                    help="write even if the groups differ in format; the reason is recorded in the manifest")
    a = ap.parse_args(argv)

    raw_root = a.raw_root or (Path(os.environ["EHR_DATA_ROOT"]) if os.environ.get("EHR_DATA_ROOT") else None)
    if raw_root is None:
        ap.error("--raw-root or EHR_DATA_ROOT is required: coverage is measured against the partition's own patients")
    dirs = partition_dirs(a.dataset_yaml)
    unknown = set(GROUPS) - set(dirs)
    if unknown:
        raise SystemExit(f"groups {sorted(unknown)} are not partitions of {a.dataset_yaml}")

    followup_path = a.all_kinds_root / FOLLOWUP_FILE
    inputs: dict[str, dict[str, Any]] = {}
    read_paths: set[Path] = set()

    def record_input(path: Path, **extra: Any) -> None:
        key = str(path.relative_to(a.all_kinds_root)) if a.all_kinds_root in path.parents else str(path)
        entry = inputs.setdefault(key, {"sha256": sha256_file(path), "bytes": path.stat().st_size, "sheets": {}})
        entry["sheets"].update(extra)
        read_paths.add(path)

    # -- the partitions' own patients --------------------------------------------------
    mrns: dict[str, set[str]] = {}
    for part in GROUPS:
        mrns[part], book = partition_mrns(raw_root, dirs[part])
        record_input(book, **{PARTITION_MRN_SHEET: {"partition": part, "patients": len(mrns[part])}})
        print(f"{part:8} partition patients {len(mrns[part]):>7,}")

    # -- read and fingerprint ----------------------------------------------------------
    tables: dict[str, dict[str, tuple[list[str], list[list[object]]]]] = {"followup": {}, "icu_transfers": {}}
    prints: dict[str, dict[str, dict[str, Any]]] = {"followup": {}, "icu_transfers": {}}
    for part, where in GROUPS.items():
        header, rows = read_sheet(followup_path, where["followup_sheet"])
        record_input(followup_path, **{where["followup_sheet"]: {"partition": part, "rows": len(rows)}})
        cols = list(FOLLOWUP_COLUMNS)
        data = project(header, rows, cols)
        for r in data:
            r.append(FOLLOWUP_CODE)
        cols.append(FOLLOWUP_CODE_COLUMN)
        data = sample_rows(data, cols.index(MRN_COLUMN), a.sample)
        tables["followup"][part] = (cols, data)
        prints["followup"][part] = fingerprint(cols, data, partition_mrns=mrns[part], mrn_column=MRN_COLUMN,
                                               date_columns=DATE_COLUMNS["followup"])
        print(f"{part:8} followup      {len(rows):>9,} rows read, {len(data):>9,} kept")

        icu_path = a.all_kinds_root / where["icu_file"]
        header, rows = read_sheet(icu_path)
        record_input(icu_path, **{"<first sheet>": {"partition": part, "rows": len(rows)}})
        data = project(header, rows, ICU_COLUMNS)
        data = sample_rows(data, ICU_COLUMNS.index(MRN_COLUMN), a.sample)
        tables["icu_transfers"][part] = (list(ICU_COLUMNS), data)
        prints["icu_transfers"][part] = fingerprint(ICU_COLUMNS, data, partition_mrns=mrns[part],
                                                    mrn_column=MRN_COLUMN, date_columns=DATE_COLUMNS["icu_transfers"])
        print(f"{part:8} icu_transfers {len(rows):>9,} rows read, {len(data):>9,} kept")

    comparison = {name: compare(prints[name]) for name in tables}
    manifest: dict[str, Any] = {
        "tool": "tools/prepare_ctpe.py",
        "all_kinds_root": str(a.all_kinds_root),
        "raw_root": str(raw_root),
        "sample": a.sample,
        "groups": GROUPS,
        "partition_dirs": dirs,
        "inputs": dict(sorted(inputs.items())),
        "fingerprint": {name: {"groups": prints[name], "comparison": comparison[name]} for name in tables},
        "outputs": {},
        "unread_inputs": unread_inputs(a.all_kinds_root, read_paths, {w["followup_sheet"] for w in GROUPS.values()}),
        "wrote": False,
        "write_despite_difference": a.write_despite_difference,
    }
    for name, verdict in comparison.items():
        print(f"fingerprint {name:14} {verdict['verdict']}: {len(verdict['hard'])} hard, {len(verdict['soft'])} soft")
        for line in verdict["hard"]:
            print(f"   HARD {line}")
    differs = any(v["verdict"] == "differs" for v in comparison.values())
    a.out.mkdir(parents=True, exist_ok=True)
    manifest_path = a.out / "prepare_manifest.json"
    if differs and not a.write_despite_difference:
        manifest_path.write_text(json.dumps(manifest, indent=2))
        print(f"STOPPED: the groups differ in format; nothing written. Comparison in {manifest_path}")
        return 2

    # -- write ------------------------------------------------------------------------
    for name, per_part in tables.items():
        for part, (cols, data) in per_part.items():
            dest = a.out / dirs[part] / f"{name}.parquet"
            n = write_parquet(cols, data, dest)
            manifest["outputs"][str(dest.relative_to(a.out))] = {"rows": n, "sha256": sha256_file(dest), "columns": cols}
            print(f"  wrote {dest.relative_to(a.out)}  {n:>9,} rows")
    manifest["wrote"] = True
    manifest_path.write_text(json.dumps(manifest, indent=2))
    print(f"manifest: {manifest_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
