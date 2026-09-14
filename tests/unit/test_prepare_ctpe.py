"""The CTPE preparation step: projection, the group fingerprint, and the refusal to
write when the four groups differ in format (remediation plan T2.J7)."""

from __future__ import annotations

import datetime as dt
import json
from pathlib import Path

import pyarrow.parquet as pq
import pytest
from openpyxl import Workbook

from tools import prepare_ctpe as prep

# --------------------------------------------------------------------------------
# helpers
# --------------------------------------------------------------------------------


def test_masked_shape_keeps_no_digit_or_letter():
    assert prep.mask_shape("AB12345678") == "aa99999999"
    assert prep.mask_shape(dt.datetime(2021, 3, 4, 5, 6, 7)) == "9999-99-99 99:99:99"
    assert prep.mask_shape(None) == "<none>"
    assert prep.mask_shape("NULL") == "aaaa"
    assert prep.mask_shape("2:13:45:10") == "9:99:99:99"


def test_cells_render_the_way_the_converter_spells_them():
    assert prep.render_cell(None) is None
    assert prep.render_cell("NULL") == "NULL"  # the delivery's literal, declared as null in the YAML
    assert prep.render_cell(12) == "12"
    assert prep.render_cell(1.0) == "1"  # 1.0 and 1 are the same cell
    assert prep.render_cell(2.5) == "2.5"
    assert prep.render_cell(dt.datetime(2021, 3, 4, 5, 6, 7, 999)) == "2021-03-04T05:06:07"
    assert prep.render_cell(dt.date(2021, 3, 4)) == "2021-03-04T00:00:00"
    with pytest.raises(TypeError):
        prep.render_cell(object())


def test_projection_keeps_only_the_named_columns_in_order():
    header = ["b", "a", "c"]
    rows = [(1, 2, 3), (4, 5)]
    assert prep.project(header, rows, ["a", "b"]) == [[2, 1], [5, 4]]
    with pytest.raises(SystemExit):
        prep.project(header, rows, ["z"])


def test_sampling_is_deterministic_and_by_patient():
    rows = [[m, i] for i, m in enumerate(["p3", "p1", "p2", "p1", "p4"])]
    kept = prep.sample_rows(rows, 0, 2)
    # sorted patients p1 p2 p3 p4 -> positions 0 and 2 -> p1 and p3, every row of each
    assert [r[0] for r in kept] == ["p3", "p1", "p1"]
    assert prep.sample_rows(rows, 0, 0) == rows


# --------------------------------------------------------------------------------
# the fingerprint
# --------------------------------------------------------------------------------


def _rows(n: int, *, null_every: int = 0, as_text: bool = False):
    out = []
    for i in range(n):
        when: object = dt.datetime(2020, 1, 1) + dt.timedelta(days=i)
        if as_text:
            when = "NULL"
        elif null_every and i % null_every == 0:
            when = None
        out.append([f"AB{i:08d}", when, 1.0])
    return out


def test_fingerprint_records_shapes_shares_and_coverage_but_no_values():
    cols = ["PAT_MRN_ID", "EFFECTIVE_DATE_DTTM", "rn"]
    rows = _rows(10, null_every=5)
    fp = prep.fingerprint(cols, rows, partition_mrns={f"AB{i:08d}" for i in range(20)},
                          mrn_column="PAT_MRN_ID", date_columns=["EFFECTIVE_DATE_DTTM"])
    assert fp["rows"] == 10 and fp["columns"] == cols
    assert fp["per_column"]["EFFECTIVE_DATE_DTTM"]["null_share"] == 0.2
    assert fp["per_column"]["EFFECTIVE_DATE_DTTM"]["types"] == {"NoneType": 0.2, "datetime": 0.8}
    assert fp["per_column"]["EFFECTIVE_DATE_DTTM"]["years"] == {"2020": 8}
    assert fp["per_column"]["PAT_MRN_ID"]["shapes"] == {"aa99999999": 1.0}
    assert fp["distinct_mrns"] == 10 and fp["coverage_of_partition"] == 0.5
    assert fp["share_of_table_mrns_outside_partition"] == 0.0
    text = json.dumps(fp)
    assert "AB0000000" not in text and "2020-01-0" not in text


def test_identical_formats_compare_as_consistent():
    cols = ["PAT_MRN_ID", "EFFECTIVE_DATE_DTTM", "rn"]
    groups = {
        g: prep.fingerprint(cols, _rows(50), partition_mrns={f"AB{i:08d}" for i in range(50)},
                            mrn_column="PAT_MRN_ID", date_columns=["EFFECTIVE_DATE_DTTM"])
        for g in ("a", "b", "c", "d")
    }
    verdict = prep.compare(groups)
    assert verdict["verdict"] == "consistent" and verdict["hard"] == []


def test_a_type_only_one_group_carries_is_a_hard_difference():
    cols = ["PAT_MRN_ID", "EFFECTIVE_DATE_DTTM", "rn"]
    mrns = {f"AB{i:08d}" for i in range(50)}
    groups = {
        "has": prep.fingerprint(cols, _rows(50), partition_mrns=mrns, mrn_column="PAT_MRN_ID"),
        "no": prep.fingerprint(cols, _rows(50, as_text=True), partition_mrns=mrns, mrn_column="PAT_MRN_ID"),
    }
    verdict = prep.compare(groups)
    assert verdict["verdict"] == "differs"
    assert any("EFFECTIVE_DATE_DTTM type 'str'" in h for h in verdict["hard"])


def test_a_null_rate_gap_and_a_coverage_gap_are_hard_differences():
    cols = ["PAT_MRN_ID", "EFFECTIVE_DATE_DTTM", "rn"]
    mrns = {f"AB{i:08d}" for i in range(50)}
    groups = {
        "has": prep.fingerprint(cols, _rows(50), partition_mrns=mrns, mrn_column="PAT_MRN_ID"),
        "no": prep.fingerprint(cols, _rows(50, null_every=4), partition_mrns=mrns | {"X"} , mrn_column="PAT_MRN_ID"),
    }
    verdict = prep.compare(groups)
    assert any("null share" in h for h in verdict["hard"])
    assert any("coverage of partition" in h for h in verdict["hard"] + verdict["soft"])


def test_different_columns_are_a_hard_difference():
    a = prep.fingerprint(["x", "y"], [[1, 2]], partition_mrns=None, mrn_column="x")
    b = prep.fingerprint(["y", "x"], [[2, 1]], partition_mrns=None, mrn_column="x")
    assert prep.compare({"a": a, "b": b})["hard"]


# --------------------------------------------------------------------------------
# end to end on a synthetic delivery
# --------------------------------------------------------------------------------

FOLLOWUP_HEADER = ["PAT_MRN_ID", "EFFECTIVE_DATE_DTTM", "rn"]


def _delivery(root: Path, raw: Path, *, break_group: str | None = None) -> None:
    """A tiny All_kinds tree and a raw export with the four partition workbooks."""
    dirs = prep.partition_dirs(Path(__file__).resolve().parents[2] / "datasets" / "ctpe.yaml")
    wb = Workbook()
    wb.remove(wb.active)
    for gi, (part, where) in enumerate(prep.GROUPS.items()):
        # the same masked shape in every group: a shape that spelled the group out
        # would be exactly the leak the fingerprint exists to catch
        mrns = [f"AB{gi}{i:07d}" for i in range(6)]
        ws = wb.create_sheet(where["followup_sheet"])
        ws.append(FOLLOWUP_HEADER)
        for i, m in enumerate(mrns[:5]):
            ws.append([m, dt.datetime(2022, 1, 1 + i, 8, 0, 0), 1])
        # the partition's own patient list
        pdir = raw / dirs[part]
        pdir.mkdir(parents=True)
        pwb = Workbook()
        pwb.active.title = "Demographics"
        pwb.active.append(["MRN", "Age"])
        for m in mrns:
            pwb.active.append([m, 50])
        pwb.save(pdir / f"{part}.xlsx")
        # the ADT workbook
        iwb = Workbook()
        iws = iwb.active
        iws.append(prep.ICU_COLUMNS)
        for i, m in enumerate(mrns[:5]):
            start = dt.datetime(2022, 1, 1 + i, 9, 0, 0)
            end: object = start + dt.timedelta(hours=30)
            dept_id: object = 100 + i
            if break_group == part:
                end, dept_id = "NULL", "NULL"
            iws.append([m, dept_id, "UNIT A", "SITE", start, end, "Admission", "N", "1:06:00:00", 5000 + i])
            iws.append([m, "NULL", "UNIT A", "SITE", start + dt.timedelta(hours=30), "NULL", "Discharge", "N", "NULL", 5000 + i])
        target = root / where["icu_file"]
        target.parent.mkdir(parents=True, exist_ok=True)
        iwb.save(target)
    ws = wb.create_sheet("35_ED_Head_CT")
    ws.append(FOLLOWUP_HEADER)
    ws.append(["Z1", dt.datetime(2022, 5, 5), 1])
    target = root / prep.FOLLOWUP_FILE
    target.parent.mkdir(parents=True)
    wb.save(target)
    other = root / "Surgical Cases" / "Surgical Cases" / "29_pulmonary_embolism.xlsx"
    other.parent.mkdir(parents=True)
    other.write_bytes(b"not read")
    (root / "ICU" / "ICU" / "29_pulmonary_embolism_notes.txt").write_bytes(b"never opened")


def test_the_prepared_tree_matches_the_partition_layout_and_the_manifest_accounts_for_every_file(tmp_path: Path):
    root, raw, out = tmp_path / "all_kinds", tmp_path / "raw", tmp_path / "prepared"
    _delivery(root, raw)
    code = prep.main(["--all-kinds-root", str(root), "--raw-root", str(raw), "--out", str(out)])
    assert code == 0
    manifest = json.loads((out / "prepare_manifest.json").read_text())
    assert manifest["wrote"] is True
    dirs = prep.partition_dirs(Path(__file__).resolve().parents[2] / "datasets" / "ctpe.yaml")
    for part, d in dirs.items():
        fu = pq.read_table(out / d / "followup.parquet")
        assert fu.column_names == FOLLOWUP_HEADER + [prep.FOLLOWUP_CODE_COLUMN]
        assert fu.num_rows == 5
        assert set(fu.column(prep.FOLLOWUP_CODE_COLUMN).to_pylist()) == {prep.FOLLOWUP_CODE}
        icu = pq.read_table(out / d / "icu_transfers.parquet")
        assert icu.column_names == prep.ICU_COLUMNS and icu.num_rows == 10
        assert all(t == "string" for t in (str(f.type) for f in icu.schema))
        assert "2022-01-01T09:00:00" in icu.column("IN_DTTM").to_pylist()
        assert "NULL" in icu.column("PAT_OUT_DTTM").to_pylist()
        assert manifest["fingerprint"]["followup"]["groups"][part]["coverage_of_partition"] == round(5 / 6, 6)
    for name, verdict in manifest["fingerprint"].items():
        assert verdict["comparison"]["verdict"] == "consistent", name
    assert len(manifest["outputs"]) == 8 and all(v["sha256"] for v in manifest["outputs"].values())
    assert all(v["sha256"] for v in manifest["inputs"].values())
    unread = {u["path"]: u["reason"] for u in manifest["unread_inputs"]}
    assert "Surgical Cases/Surgical Cases/29_pulmonary_embolism.xlsx" in unread
    assert "P-J12" in unread["Surgical Cases/Surgical Cases/29_pulmonary_embolism.xlsx"]
    assert "never opened" in unread["ICU/ICU/29_pulmonary_embolism_notes.txt"]
    assert unread[prep.FOLLOWUP_FILE + "::35_ED_Head_CT"].startswith("sheet for another cohort")
    text = json.dumps(manifest)
    assert "AB00000001" not in text  # no MRN reaches the manifest


def test_a_group_whose_format_differs_stops_the_write(tmp_path: Path):
    root, raw, out = tmp_path / "all_kinds", tmp_path / "raw", tmp_path / "prepared"
    _delivery(root, raw, break_group="29_no")
    code = prep.main(["--all-kinds-root", str(root), "--raw-root", str(raw), "--out", str(out)])
    assert code == 2
    manifest = json.loads((out / "prepare_manifest.json").read_text())
    assert manifest["wrote"] is False and not manifest["outputs"]
    assert manifest["fingerprint"]["icu_transfers"]["comparison"]["verdict"] == "differs"
    assert not list(out.rglob("*.parquet"))
    # an explicit, recorded override writes anyway
    code = prep.main(["--all-kinds-root", str(root), "--raw-root", str(raw), "--out", str(out),
                      "--write-despite-difference", "test override"])
    assert code == 0
    manifest = json.loads((out / "prepare_manifest.json").read_text())
    assert manifest["wrote"] is True and manifest["write_despite_difference"] == "test override"


def test_sampling_keeps_every_nth_patient(tmp_path: Path):
    root, raw, out = tmp_path / "all_kinds", tmp_path / "raw", tmp_path / "prepared"
    _delivery(root, raw)
    assert prep.main(["--all-kinds-root", str(root), "--raw-root", str(raw), "--out", str(out), "--sample", "2"]) == 0
    manifest = json.loads((out / "prepare_manifest.json").read_text())
    assert manifest["sample"] == 2
    rows = [v["rows"] for k, v in manifest["outputs"].items() if k.endswith("followup.parquet")]
    assert rows == [3, 3, 3, 3]
